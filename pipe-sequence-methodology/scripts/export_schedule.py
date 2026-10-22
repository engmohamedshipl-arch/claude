"""Export the recalculated sequence workbook to Primavera P6 (.xer) and MS Project (.xml, MSPDI) with MPXJ.
Args: workbook.xlsx out.xer out.xml
"""
import sys, math, datetime, collections
import openpyxl
import jpype, mpxj  # noqa: F401  (mpxj puts its jars on the classpath)

WB, OUT_XER, OUT_XML = sys.argv[1:4]
wb = openpyxl.load_workbook(WB, data_only=True)

# ----------------------------------------------------------------- read the workbook
inp = {r[0]: r[2] for r in wb['Inputs'].iter_rows(min_row=5, values_only=True) if r[0] and r[2] is not None}
START, WD = inp['START'], inp['WD']
START_D = START.date() if isinstance(START, datetime.datetime) else START
D_PILEDES, D_XING = inp['D_PILEDES'], inp['D_XING']

wsK = wb['Packages']
kh = [c.value for c in wsK[5]]
pkg = {}
for row in wsK.iter_rows(min_row=6, values_only=True):
    if row[1] and str(row[1]).startswith('UG') and row[2]:
        d = dict(zip(kh, row)); pkg[d['Package']] = d

wsT = wb['Activities']
th = [c.value for c in wsT[5]]
acts, order = {}, []
for row in wsT.iter_rows(min_row=6, values_only=True):
    if not row[0]:
        continue
    d = dict(zip(th, row)); acts[d['Activity ID']] = d; order.append(d['Activity ID'])

rels = []
for row in wb['Relationships'].iter_rows(min_row=5, values_only=True):
    if row[0]:
        rels.append(dict(pred=row[0], succ=row[1], type=row[2], lag=row[3] or 0, why=row[4]))

dur = {a: int(acts[a]['Duration (work days)'] or 0) for a in acts}
def is_ms(a):
    return acts[a]['Code'] == 'A010'

# ----------------------------------------------------------------- logic: lags to working days, zero-duration bypass
def lag_wd(r):
    if r['why'].startswith('Just-in-time dewatering'):
        return -dur[r['succ']]
    return int(round(r['lag'] * WD / 7.0))

edges = []
chain_prev = {a: acts[a]['Chain predecessor'] for a in acts if acts[a]['Chain predecessor']}
def live_chain_pred(a):
    p = chain_prev.get(a)
    while p and dur[p] == 0:
        p = chain_prev.get(p)
    return p
for r in rels:
    if r['why'].startswith('Crew continuity'):
        continue
    edges.append((r['pred'], r['succ'], r['type'], lag_wd(r)))
for a in acts:
    if a in chain_prev and dur[a] > 0:
        p = live_chain_pred(a)
        if p:
            edges.append((p, a, 'FS', 0))

zero = [a for a in order if dur[a] == 0 and not is_ms(a)]
for z in zero:
    ins = [e for e in edges if e[1] == z]
    outs = [e for e in edges if e[0] == z]
    edges = [e for e in edges if z not in (e[0], e[1])]
    for p, _, t1, l1 in ins:
        for _, s, t2, l2 in outs:
            if t1[1] == 'F' and t2[0] == 'S':      # finish-side input cannot drive a start-side output
                continue
            edges.append((p, s, t1[0] + t2[1], l1 + l2))
# dedupe: same pred/succ/type keeps the largest lag
best = {}
for p, s, t, l in edges:
    k = (p, s, t)
    best[k] = max(best.get(k, -10**9), l)
edges = [(p, s, t, l) for (p, s, t), l in best.items()]
keep = [a for a in order if a not in set(zero)]
print('activities', len(keep), 'removed zero-duration', len(zero), 'relationships', len(edges))

# MS Project allows one link per pair: keep the link that drives the workbook dates
def driving_type(p, s, types):
    if len(types) == 1:
        return types[0]
    A, B = acts[p], acts[s]
    if 'FF' in types and 'SS' in types:
        ff_lag = B['FF lag (cal. days)'] or 0
        ff_drives = B['Finish'] == A['Finish'] + datetime.timedelta(days=ff_lag) and B['Finish'] > B['Start'] + datetime.timedelta(days=B['Duration (cal. days)'] or 0)
        return 'FF' if ff_drives else 'SS'
    for t in ('FS', 'SS', 'FF', 'SF'):
        if t in types:
            return t
pairs = collections.defaultdict(dict)
for p, s, t, l in edges:
    pairs[(p, s)][t] = l
edges_msp = [(p, s, (t := driving_type(p, s, list(ts))), ts[t]) for (p, s), ts in pairs.items()]

# ----------------------------------------------------------------- build with MPXJ
jpype.startJVM()
from org.mpxj import (ProjectFile, Duration, TimeUnit, RelationType, Relation, ConstraintType, LocalTimeRange,
                      ActivityType, TaskType)
from org.mpxj.writer import UniversalProjectWriter, FileFormat
from org.mpxj.reader import UniversalProjectReader
from java.time import LocalDateTime, LocalTime, DayOfWeek

RT = {'FS': RelationType.FINISH_START, 'SS': RelationType.START_START, 'FF': RelationType.FINISH_FINISH, 'SF': RelationType.START_FINISH}
def ldt(d, h=8):
    return LocalDateTime.of(d.year, d.month, d.day, h, 0)

def build(edge_list, for_p6):
    pf = ProjectFile()
    cal = pf.addDefaultBaseCalendar()
    cal.setName('AW 6-day Sat-Thu 8h')
    for day in (DayOfWeek.SATURDAY, DayOfWeek.SUNDAY, DayOfWeek.MONDAY, DayOfWeek.TUESDAY, DayOfWeek.WEDNESDAY, DayOfWeek.THURSDAY):
        cal.setWorkingDay(day, True)
        h = cal.getCalendarHours(day) or cal.addCalendarHours(day)
        h.clear()
        h.add(LocalTimeRange(LocalTime.of(8, 0), LocalTime.of(12, 0)))
        h.add(LocalTimeRange(LocalTime.of(13, 0), LocalTime.of(17, 0)))
    cal.setWorkingDay(DayOfWeek.FRIDAY, False)
    h = cal.getCalendarHours(DayOfWeek.FRIDAY)
    if h is not None:
        h.clear()
    pf.setDefaultCalendar(cal)
    pp = pf.getProjectProperties()
    pp.setName('Alex West WWTP - Underground pipe works (201 packages)')
    pp.setProjectID('AW-PIPE-SEQ')
    pp.setStartDate(ldt(START))
    pp.setStatusDate(ldt(START))
    pp.setMinutesPerDay(jpype.JInt(480)); pp.setMinutesPerWeek(jpype.JInt(480 * WD)); pp.setDaysPerMonth(jpype.JInt(26))
    pp.setDefaultCalendar(cal) if hasattr(pp, 'setDefaultCalendar') else None

    def nxt_work(d):
        while d.weekday() == 4:   # Friday
            d += datetime.timedelta(days=1)
        return d

    root = pf.addTask(); root.setName('Alex West underground pipe works'); root.setActivityID('AW-PIPE'); root.setWBS('AW-PIPE')
    fronts, lines_, pkgs_, tasks = {}, {}, {}, {}
    fdesc = {r[0]: r[1] for r in wb['Fronts'].iter_rows(min_row=5, max_row=13, values_only=True) if r[0]}
    for a in keep:
        A = acts[a]; f, l, p = A['Front'], A['Line'], A['Package']
        if f not in fronts:
            t = root.addTask(); t.setName(f'{f} {fdesc.get(f, "")}'); t.setActivityID(f); t.setWBS(f'AW-PIPE.{f}'); fronts[f] = t
        if (f, l) not in lines_:
            t = fronts[f].addTask(); t.setName(l); t.setActivityID(f'{f}.{l}'); t.setWBS(f'AW-PIPE.{f}.{l}'); lines_[(f, l)] = t
        if p not in pkgs_:
            P = pkg[p]
            t = lines_[(f, l)].addTask(); t.setName(f"{p} | {P['Planning name']} | {P['Package type']} | {P['Excavation method']}")
            t.setActivityID(p); t.setWBS(f'AW-PIPE.{f}.{l}.{p}'); pkgs_[p] = t
        t = pkgs_[p].addTask()
        t.setName(A['Activity'] if for_p6 else f"{a} {A['Activity']}"); t.setActivityID(a); t.setCalendar(cal)
        t.setType(TaskType.FIXED_DURATION)
        d = dur[a]
        if is_ms(a):
            t.setMilestone(True); t.setActivityType(ActivityType.START_MILESTONE)
        else:
            t.setActivityType(ActivityType.TASK_DEPENDENT)
        D = Duration.getInstance(d, TimeUnit.DAYS)
        t.setDuration(D); t.setRemainingDuration(D)
        st = ldt(nxt_work(A['Start'].date() if isinstance(A['Start'], datetime.datetime) else A['Start']))
        fin = st if d == 0 else cal.getDate(st, D)
        t.setStart(st); t.setFinish(fin); t.setEarlyStart(st); t.setEarlyFinish(fin)
        if for_p6:
            t.setPlannedStart(st); t.setPlannedFinish(fin)
        # gates as start-no-earlier-than constraints
        gate = None
        if A['Code'] == 'A010':
            gate = A['Start']
        elif A['Code'] == 'A030' and D_PILEDES > START:
            gate = D_PILEDES
        elif A['Code'] == 'A060' and pkg[p]['Crossing ID hold'] == 'Yes' and D_XING > START:
            gate = D_XING
        if gate is not None:
            g = gate.date() if isinstance(gate, datetime.datetime) else gate
            t.setConstraintType(ConstraintType.START_NO_EARLIER_THAN); t.setConstraintDate(ldt(nxt_work(g)))
        qty = A['Quantity']
        note = f"Front {f} | Line {l} | Package {p}"
        if qty not in (None, ''):
            note += f" | Quantity {qty:,.0f} {A['Unit']}"
        if A['Other predecessors (type)']:
            note += f" | Workbook logic: {A['Other predecessors (type)']}"
        t.setNotes(note)
        t.setText(1, f); t.setText(2, l); t.setText(3, p); t.setText(4, A['Code'])
        if qty not in (None, ''):
            t.setNumber(1, float(qty))
        tasks[a] = t
    for p, s, typ, l in edge_list:
        tasks[s].addPredecessor(Relation.Builder().predecessorTask(tasks[p]).type(RT[typ]).lag(Duration.getInstance(l, TimeUnit.DAYS)))
    return pf

p6 = build(edges, True)
UniversalProjectWriter(FileFormat.XER).write(p6, OUT_XER)
msp = build(edges_msp, False)
UniversalProjectWriter(FileFormat.MSPDI).write(msp, OUT_XML)

# ----------------------------------------------------------------- read back both files
for fn in (OUT_XER, OUT_XML):
    q = UniversalProjectReader().read(fn)
    ts = [t for t in q.getTasks() if t.getText(4) and str(t.getText(4)).startswith('A') or (t.getActivityID() and str(t.getActivityID()).startswith('UG') and '-A' in str(t.getActivityID()))]
    npred = sum(t.getPredecessors().size() for t in q.getTasks())
    ncon = sum(1 for t in ts if t.getConstraintType() is not None and str(t.getConstraintType()) == 'START_NO_EARLIER_THAN')
    cal = q.getDefaultCalendar()
    print(fn.split('/')[-1], '| activities', len(ts), '| relationships', npred, '| SNET constraints', ncon,
          '| calendar', cal.getName() if cal else None, '| Fri working', cal.isWorkingDay(DayOfWeek.FRIDAY) if cal else None,
          '| Sat working', cal.isWorkingDay(DayOfWeek.SATURDAY) if cal else None)

# ----------------------------------------------------------------- forward pass on the exported logic (6-day calendar)
def wd_index(d):
    d = d.date() if isinstance(d, datetime.datetime) else d
    n, x = 0, START_D
    step = 1 if d >= START_D else -1
    while x != d:
        x += datetime.timedelta(days=step)
        if x.weekday() != 4:
            n += step
    return n
def from_index(n):
    x, k = START_D, 0
    while k < n:
        x += datetime.timedelta(days=1)
        if x.weekday() != 4:
            k += 1
    return x
for label, el in (('P6', edges), ('MSP', edges_msp)):
    preds = collections.defaultdict(list); succs = collections.defaultdict(list); indeg = collections.Counter({a: 0 for a in keep})
    for p, s_, t, l in el:
        preds[s_].append((p, t, l)); succs[p].append(s_); indeg[s_] += 1
    cons = {}
    for a in keep:
        A = acts[a]
        if A['Code'] == 'A010': cons[a] = wd_index(A['Start'])
    q = [a for a in keep if indeg[a] == 0]; ES, EF = {}, {}; seen = 0
    while q:
        a = q.pop(); seen += 1
        es = max(0, cons.get(a, 0))
        for p, t, l in preds[a]:
            if t == 'FS': es = max(es, EF[p] + l)
            elif t == 'SS': es = max(es, ES[p] + l)
            elif t == 'FF': es = max(es, EF[p] + l - dur[a])
            elif t == 'SF': es = max(es, ES[p] + l - dur[a])
        ES[a] = es; EF[a] = es + dur[a]
        for s_ in succs[a]:
            indeg[s_] -= 1
            if indeg[s_] == 0: q.append(s_)
    fin = from_index(max(EF.values()))
    print(label, 'forward pass: scheduled', seen, 'of', len(keep), '(loops if fewer) | project finish', fin)
