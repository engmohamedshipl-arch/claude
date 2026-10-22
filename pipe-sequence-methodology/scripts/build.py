"""Build Alex West pipe construction sequence & methodology workbook.
Args: package_workbook.xlsx aer_piles.json aer_weekly.json manifest.json out.xlsx
"""
import sys, json, re, heapq, datetime, collections
import numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter as CL
from openpyxl.comments import Comment

SRC, PILES, WEEKLY, MANIFEST, OUT = sys.argv[1:6]

# ----------------------------------------------------------------- source data
seg = pd.read_excel(SRC, 'Segments', header=3)
seg = seg[seg['Segment ID'].astype(str).str.contains('-S')].reset_index(drop=True)
pcs = pd.read_excel(SRC, 'Pieces', header=3)
pcs = pcs[pcs['Line'].astype(str).str.startswith('UG')].reset_index(drop=True)
lines = pd.read_excel(SRC, 'Lines', header=3)
lines = lines[lines['Line'].astype(str).str.startswith('UG')].reset_index(drop=True)
aer = pd.DataFrame(json.load(open(PILES)))
aer['E'] = aer.x + 774450.0
aer['N'] = aer.y + 3450600.0
weekly = json.load(open(WEEKLY))
manifest = {a['id']: a for a in json.load(open(MANIFEST))['assets']}

NG, GWL, SLAB = -0.7, -1.5, 0.5   # only used for the geometric interaction test

FRONTS = [
    ('F1', 'Headworks & primary (GRT-PST-IPS)', ['UG-01', 'UG-02', 'UG-03', 'UG-04', 'UG-24'], 1),
    ('F2', 'Bypass corridor: IPS bypass and outfall twin', ['UG-05', 'UG-06', 'UG-25', 'UG-26'], 1),
    ('F3', 'Distribution & FCT risers', ['UG-08', 'UG-09', 'UG-10', 'UG-14E', 'UG-14F', 'UG-14G', 'UG-14H',
                                        'UG-16A', 'UG-16B', 'UG-16C', 'UG-16D', 'UG-27I', 'UG-27J', 'UG-27K', 'UG-27L'], 1),
    ('F4', 'FCT effluent A-H to CCT', ['UG-07', 'UG-07A', 'UG-15', 'UG-15A'], 1),
    ('F5', 'FCT effluent I-L to CCT', ['UG-12', 'UG-12A'], 1),
    ('F6', 'Sludge A-D,G,H to RAS/WAS + UG-23', ['UG-11', 'UG-11A', 'UG-11B', 'UG-23'], 1),
    ('F7', 'Sludge I-L,E,F to RAS/WAS', ['UG-13', 'UG-13A', 'UG-13B'], 1),
    ('F8', 'Effluent, outfall & CCT links', ['UG-17', 'UG-18', 'UG-19', 'UG-20', 'UG-21', 'UG-22', 'UG-28', 'UG-28A'], 1),
]
line_front = {l: f for f, _, ls, _ in FRONTS for l in ls}
assert set(line_front) == set(lines['Line']), set(lines['Line']) ^ set(line_front)

NON_STRUCT = re.compile(r'^(EL\d+|CH\d+|TEE\d*)$')
NOT_MODELLED = {'EXPST', 'BYPCH', 'CHUG25', 'OUTF', 'TEPS'}
ASSET_OF = {'GRT': 'AW-GRIT', 'PST': 'AW-PST', 'PST08': 'AW-PST', 'IPS': 'AW-IPS', 'CCT': 'AW-CCT',
            'SDCA': 'AW-SDC-A', 'SDCB': 'AW-SDC-B', 'SDCC': 'AW-SDC-C', 'RASPS': 'AW-RASWAS',
            'AER': 'AW-AER', 'WFT': 'AW-FIRE'}
for c in 'ABCDEFGHIJKL':
    ASSET_OF['FST' + c] = 'AW-FST-' + c

# ----------------------------------------------------------------- packages
pk = seg[['Line', 'Segment ID', 'Path', 'Planning name', 'Pipe type', 'Printed size', 'From (short)', 'To (short)',
          'From', 'To', 'Trench / corridor type', 'R at start (m)', 'R at end (m)', 'Fittings',
          'Crossings and parallel lines', 'Holds before release', 'Depends on / sequence', 'Direction basis']].copy()
pk.columns = ['line', 'id', 'path', 'pname', 'ptype', 'size', 'fs', 'ts', 'ffull', 'tfull', 'trench', 'rs', 're',
              'fittings', 'cross', 'holds', 'depends', 'dirbasis']
pk['front'] = pk.line.map(line_front)
pk['kind'] = np.where(pk.trench.str.startswith('Short connection'), 'Short connection',
                      np.where(pk.trench.str.startswith('Tie-in'), 'Tie-in', 'Run'))
structs = {}
def ends(r):
    out = []
    for code, full in ((r.fs, r.ffull), (r.ts, r.tfull)):
        if not NON_STRUCT.match(str(code)):
            out.append(code)
            structs.setdefault(code, full)
    return out if r.kind != 'Run' else []
pk['struct'] = pk.apply(ends, axis=1)
for r in pk.itertuples():          # collect structure names for run ends too (for the register)
    for code, full in ((r.fs, r.ffull), (r.ts, r.tfull)):
        if not NON_STRUCT.match(str(code)):
            structs.setdefault(code, full)
pk['nfit'] = pk.fittings.apply(lambda s: sum(int(x) for x in re.findall(r'(\d+)\s*x', str(s))))
pk['xing'] = pk.cross.str.contains('UNASSIGNED', na=False)
pk['seq'] = pk.groupby('path').cumcount()
pk['npath'] = pk.groupby('path')['id'].transform('size')
idx = {p: i for i, p in enumerate(pk.id)}

# aeration piles (pending/defect/unknown) within 15 m of any package piece
def segdist(px, py, ax, ay, bx, by):
    dx, dy = bx - ax, by - ay
    L2 = dx * dx + dy * dy
    t = np.clip(((px - ax) * dx + (py - ay) * dy) / (L2 if L2 > 0 else 1), 0, 1)
    return np.hypot(px - (ax + t * dx), py - (ay + t * dy))
open_p = aer[aer.status != 'done']
zones, zcnt = [], []
for pid in pk.id:
    g = pcs[pcs.Package == pid]
    d = np.full(len(open_p), 1e9)
    for _, a in g.iterrows():
        d = np.minimum(d, segdist(open_p.E.values, open_p.N.values, a['Start E'], a['Start N'], a['End E'], a['End N']))
    w = open_p[d <= 15.0]
    zones.append(sorted(set(int(z) for z in w.zone)))
    zcnt.append(len(w))
pk['aerz'] = zones
pk['aern'] = zcnt

# geometric interactions between packages of different lines (axis within 8 m)
pts = []
for _, r in pcs.iterrows():
    n = max(2, int(np.ceil(r['Plan length (m)'])) + 1)
    for t in np.linspace(0, 1, n):
        cl = r['Start CL (m)'] + (r['End CL (m)'] - r['Start CL (m)']) * t
        pts.append((r['Package'], r['Line'], r['Start E'] + (r['End E'] - r['Start E']) * t,
                    r['Start N'] + (r['End N'] - r['Start N']) * t, cl - r['Planning dia (m)'] / 2 - SLAB))
P = pd.DataFrame(pts, columns=['pkg', 'line', 'E', 'N', 'R'])
XY = P[['E', 'N']].values
pairs = []
for i0 in range(0, len(XY), 1000):
    blk = XY[i0:i0 + 1000]
    dd = np.hypot(blk[:, None, 0] - XY[None, :, 0], blk[:, None, 1] - XY[None, :, 1])
    ii, jj = np.nonzero(dd <= 8.0)
    ii = ii + i0
    k = jj > ii
    pairs.append(np.stack([ii[k], jj[k]], 1))
pairs = np.concatenate(pairs)
a = P.iloc[pairs[:, 0]].reset_index(drop=True)
b = P.iloc[pairs[:, 1]].reset_index(drop=True)
m = a.line.values != b.line.values
a, b = a[m].reset_index(drop=True), b[m].reset_index(drop=True)
sw = a.pkg.values > b.pkg.values
p1 = np.where(sw, b.pkg.values, a.pkg.values); p2 = np.where(sw, a.pkg.values, b.pkg.values)
r1 = np.where(sw, b.R.values, a.R.values); r2 = np.where(sw, a.R.values, b.R.values)
it = pd.DataFrame({'p1': p1, 'p2': p2, 'dr': r1 - r2}).groupby(['p1', 'p2'])['dr'].median().reset_index()
shared, deep = [], []      # shared: (leader, follower); deep: (deeper, shallower)
for r in it.itertuples():
    if abs(r.dr) <= 1.0:
        l1, l2 = idx[r.p1], idx[r.p2]
        lead, fol = (r.p1, r.p2) if r.dr <= 0 else (r.p2, r.p1)
        shared.append((lead, fol))
    else:
        deep.append((r.p1, r.p2) if r.dr < 0 else (r.p2, r.p1))

# closure: tie-in lays after the adjacent run package on its path
closure = {}
for r in pk.itertuples():
    if r.kind != 'Tie-in':
        continue
    same = pk[(pk.path == r.path) & (pk.kind == 'Run')]
    if len(same) == 0:
        continue
    nb = same.iloc[(same.seq - r.seq).abs().argsort()].iloc[0]
    closure[r.id] = nb.id


cent = {p: np.array([g[['Start E', 'End E']].values.mean(), g[['Start N', 'End N']].values.mean()]) for p, g in pcs.groupby('Package')}
# ----------------------------------------------------------------- global order (topological)
line_depth = lines.set_index('Line')['R deepest (m)'].to_dict()
def key(pid):
    r = pk.iloc[idx[pid]]
    lay = r.npath - 1 - r.seq          # downstream first
    return (line_depth[r.line], r.line, r.path, lay)
# shared corridors that cross fronts go to the front holding most of their length (one crew per corridor)
plen = dict(zip(seg['Segment ID'], seg['Plan length (m)']))
cpar = {p: p for p in pk.id}
def croot(x):
    while cpar[x] != x:
        cpar[x] = cpar[cpar[x]]; x = cpar[x]
    return x
for l_, f_ in shared:
    cpar[croot(l_)] = croot(f_)
comps = collections.defaultdict(list)
for p in pk.id:
    comps[croot(p)].append(p)
reassigned = []
for ms in comps.values():
    fl = collections.Counter()
    for p in ms:
        fl[pk.at[idx[p], 'front']] += plen[p]
    if len(fl) > 1:
        best = fl.most_common(1)[0][0]
        for p in ms:
            if pk.at[idx[p], 'front'] != best:
                reassigned.append((p, pk.at[idx[p], 'front'], best)); pk.at[idx[p], 'front'] = best
print('reassigned', len(reassigned))
FRONTS = [fr_ for fr_ in FRONTS if (pk.front == fr_[0]).any()]
dag_edges = [(x, y, 'deep') for x, y in deep] + [(x, y, 'shared') for x, y in shared] + [(nb, t, 'closure') for t, nb in closure.items()]
shared_set = set(shared)
# contract same-front shared-trench components into one unit so they are built back to back
par = {p: p for p in pk.id}
def root(x):
    while par[x] != x:
        par[x] = par[par[x]]; x = par[x]
    return x
for l_, f_ in shared:
    if pk.iloc[idx[l_]].front == pk.iloc[idx[f_]].front:
        par[root(l_)] = root(f_)
unit = {p: root(p) for p in pk.id}
members = collections.defaultdict(list)
for p in pk.id:
    members[unit[p]].append(p)
ukey = {u: min(key(p) for p in ms) for u, ms in members.items()}
usucc = collections.defaultdict(set); uindeg = collections.Counter({u: 0 for u in members})
for x, y, _ in dag_edges:
    ux, uy = unit[x], unit[y]
    if ux != uy and uy not in usucc[ux]:
        usucc[ux].add(uy); uindeg[uy] += 1
heap = [(ukey[u], u) for u in members if uindeg[u] == 0]; heapq.heapify(heap)
uorder = []; done_u = set()
while len(uorder) < len(members):
    if not heap:
        rest = [u for u in members if u not in done_u]
        u = min(rest, key=lambda u: ukey[u]); uindeg[u] = 0; heapq.heappush(heap, (ukey[u], u))
        print('unit cycle broken at', u)
    _, u = heapq.heappop(heap)
    if u in done_u:
        continue
    done_u.add(u); uorder.append(u)
    for v in usucc[u]:
        uindeg[v] -= 1
        if uindeg[v] == 0:
            heapq.heappush(heap, (ukey[v], v))
# expand: inside a unit, topological on internal edges, ties by key
order = []
for u in uorder:
    ms = set(members[u])
    ins = collections.Counter({p: 0 for p in ms}); out = collections.defaultdict(list)
    for x, y, _ in dag_edges:
        if x in ms and y in ms:
            out[x].append(y); ins[y] += 1
    got = []; placed = set()
    while len(got) < len(ms):
        avail = [q for q in ms if q not in placed and ins[q] <= 0]
        if not avail:
            q = min((q for q in ms if q not in placed), key=key); ins[q] = 0; avail = [q]
        lastc = cent[got[-1]] if got else None
        def score(q):
            links = sum(1 for o in placed if (o, q) in shared_set or (q, o) in shared_set)
            dist = np.hypot(*(cent[q] - lastc)) if lastc is not None else 0.0
            return (-links, dist, key(q))
        q = min(avail, key=score)
        got.append(q); placed.add(q)
        for r2 in out[q]:
            ins[r2] -= 1
    order += got
rank = {p: i for i, p in enumerate(order)}
pk['grank'] = pk.id.map(rank)

chains = {}   # front -> ordered run packages
for f, *_ in FRONTS:
    chains[f] = [p for p in order if pk.iloc[idx[p]].front == f and pk.iloc[idx[p]].kind == 'Run']
chain_prev = {}
for f, lst in chains.items():
    for i in range(1, len(lst)):
        chain_prev[lst[i]] = lst[i - 1]

deep_preds = collections.defaultdict(list); ss_preds = collections.defaultdict(list); bf_preds = collections.defaultdict(list)
for d_, s_ in deep:
    deep_preds[s_].append(d_)
for l_, f_ in shared:
    ss_preds[f_].append(l_)
    bf_preds[l_].append(f_)
    bf_preds[f_].append(l_)

# ----------------------------------------------------------------- activity graph cycle check
ACTS = [('A010', 'Release: drawings, holds, access and gates'),
        ('A020', 'Set-out, crossing survey and working platform'),
        ('A030', 'Pipe-support piling from platform (incl. integrity tests)'),
        ('A040', 'Install sheet piles (shored lengths)'),
        ('A050', 'Install dewatering and draw down'),
        ('A060', 'Excavate and strut to formation'),
        ('A070', 'Trim piles, blinding and piled slab (incl. cure)'),
        ('A080', 'Lay and joint pipe (incl. fittings, tie-in)'),
        ('A090', 'Surround, joint test and line test'),
        ('A100', 'Backfill to NG and remove struts'),
        ('A110', 'Dewatering off, extract sheets, reinstate')]
CODES = [c for c, _ in ACTS]
CHAINED = {'A030', 'A040', 'A060', 'A070', 'A080', 'A100'}
edges = []
for p in pk.id:
    for i in range(1, len(CODES)):
        edges.append(((p, CODES[i - 1]), (p, CODES[i])))
    if p in chain_prev:
        for c in CHAINED:
            edges.append(((chain_prev[p], c), (p, c)))
    for d_ in deep_preds[p]:
        edges.append(((d_, 'A100'), (p, 'A060')))
    for l_ in ss_preds[p]:
        edges.append(((l_, 'A060'), (p, 'A060')))
    for o_ in bf_preds[p]:
        edges.append(((o_, 'A080'), (p, 'A100')))
    if p in closure:
        edges.append(((closure[p], 'A080'), (p, 'A080')))
        edges.append(((closure[p], 'A030'), (p, 'A020')))
    if p in chain_prev:
        edges.append(((chain_prev[p], 'A030'), (p, 'A020')))
        edges.append(((chain_prev[p], 'A060'), (p, 'A050')))
def find_cycle(edges):
    g = collections.defaultdict(list)
    for u, v in edges:
        g[u].append(v)
    color = {}
    for s0 in list(g):
        if color.get(s0):
            continue
        stack = [(s0, iter(g[s0]))]; color[s0] = 1; path = [s0]
        while stack:
            n, itr = stack[-1]
            nxt = next(itr, None)
            if nxt is None:
                color[n] = 2; stack.pop(); path.pop(); continue
            c = color.get(nxt, 0)
            if c == 1:
                cyc = path[path.index(nxt):] + [nxt]
                return list(zip(cyc[:-1], cyc[1:]))
            if c == 0:
                color[nxt] = 1; stack.append((nxt, iter(g[nxt]))); path.append(nxt)
    return []
dropped = []
cyc = find_cycle(edges)
while cyc:
    ce = set(cyc)
    cl = [(u, v) for u, v in cyc if u[1] == 'A080' and v[1] == 'A080' and v[0] in closure and closure[v[0]] == u[0]]
    dp = [(u, v) for u, v in cyc if u[1] == 'A100' and v[1] == 'A060' and u[0] in deep_preds[v[0]]]
    bf = [(u, v) for u, v in cyc if u[1] == 'A080' and v[1] == 'A100' and u[0] in bf_preds[v[0]]]
    if bf:
        u, v = bf[0]; bf_preds[v[0]].remove(u[0]); dropped.append(('shared-backfill', u[0], v[0]))
    elif cl:
        u, v = cl[0]; closure.pop(v[0]); dropped.append(('closure', v[0], u[0]))
    elif dp:
        u, v = dp[0]; deep_preds[v[0]].remove(u[0]); dropped.append(('deep', u[0], v[0]))
    else:
        raise SystemExit('unresolvable cycle')
    edges.remove((u, v))
    cyc = find_cycle(edges)
print('interactions: shared', len(shared), 'deep-first', len(deep), 'dropped', dropped)

# ----------------------------------------------------------------- workbook styles
wb = Workbook()
F = 'Arial'
fnt = Font(name=F, size=10)
bold = Font(name=F, size=10, bold=True)
title = Font(name=F, size=14, bold=True, color='1F3864')
sub = Font(name=F, size=10, italic=True, color='595959')
hdr_font = Font(name=F, size=10, bold=True, color='FFFFFF')
blue = Font(name=F, size=10, color='0000FF')
green = Font(name=F, size=10, color='008000')
hdr_fill = PatternFill('solid', fgColor='1F3864')
sec_fill = PatternFill('solid', fgColor='D9E1F2')
yel = PatternFill('solid', fgColor='FFFF00')
grey = PatternFill('solid', fgColor='F2F2F2')
thin = Side(style='thin', color='BFBFBF')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical='top')
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
DATE = 'dd-mmm-yy'

def style_range(ws, rng, font=None, fill=None, fmt=None, align=None, border=True):
    for row in ws[rng]:
        for c in row:
            if font: c.font = font
            if fill: c.fill = fill
            if fmt: c.number_format = fmt
            if align: c.alignment = align
            if border: c.border = box

def header(ws, row, cols, widths=None):
    for j, h in enumerate(cols, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = hdr_font; c.fill = hdr_fill; c.alignment = center; c.border = box
    ws.row_dimensions[row].height = 42
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[CL(j)].width = w

def put(ws, r, c, v, font=fnt, fmt=None, fill=None, align=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font; cell.border = box
    if fmt: cell.number_format = fmt
    if fill: cell.fill = fill
    if align: cell.alignment = align
    return cell

def name(n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=ref)

# ----------------------------------------------------------------- Read Me (filled at end)
wsR = wb.active; wsR.title = 'Read Me'

# ----------------------------------------------------------------- Inputs
wsI = wb.create_sheet('Inputs')
wsI['A1'] = 'Inputs - levels, ground, groundwater, piles, programme gates and productivity'; wsI['A1'].font = title
wsI['A2'] = 'Yellow = edit. Every quantity, method flag, duration and date in this workbook recalculates from these cells.'; wsI['A2'].font = sub
header(wsI, 4, ['Name', 'Description', 'Value', 'Unit', 'Source / status'], [14, 62, 13, 12, 90])
inputs = [
    ('S', 'Levels and ground'),
    ('NG', 'Adopted natural ground level', -0.7, 'm', 'Package workbook Inputs (D-004, Mohamed 20 Oct 2026). Single site value.'),
    ('FL', 'Finished ground level', 1.0, 'm', 'Package workbook Inputs (D-004; DLU sections FGL +1.00).'),
    ('GWL', 'Groundwater level for design of temporary works', -1.5, 'm', 'Mohamed, 22 Oct 2026.'),
    ('SLAB', 'Piled slab + blinding below external pipe bottom (excavation bottom R = pipe bottom - SLAB)', 0.5, 'm', 'D-001, user-stated basis. All pipes on piles and a piled slab (Mohamed, 22 Oct 2026). No pile/slab design yet.'),
    ('SLABOVH', 'Piled slab overhang beyond pipe outside, each side', 0.5, 'm', 'ASSUMPTION for concrete quantity only.'),
    ('WS', 'Working space, pipe outside to trench toe / sheet line', 1.0, 'm', 'AGREED by Mohamed, 21 Oct 2026.'),
    ('SLOPE_A', 'Open-cut side slope above GWL (H per 1 V)', 1.5, 'H:V', 'ASSUMPTION (very soft silty clay). Package workbook value.'),
    ('SLOPE_B', 'Open-cut side slope below GWL, dewatered (H per 1 V)', 2.0, 'H:V', 'ASSUMPTION: flatter below water in soft clay. Geotechnical confirmation needed.'),
    ('MAXOC', 'Maximum excavation depth below GWL allowed in open cut (deeper = sheet piles regardless of space)', 2.5, 'm', 'ASSUMPTION: base heave / slope stability limit in soft clay below water. See Sensitivity sheet.'),
    ('DD_BELOW', 'Drawdown target below excavation formation', 0.5, 'm', 'Planning rule of thumb.'),
    ('WP_LIMIT', 'Maximum drawdown for single-stage wellpoints', 4.5, 'm', 'Planning rule of thumb; deeper needs deep wells / two stages.'),
    ('GAM_D', 'Unit weight of backfill above GWL', 18.0, 'kN/m3', 'ASSUMPTION.'),
    ('GAM_S', 'Saturated unit weight of backfill below GWL', 18.0, 'kN/m3', 'ASSUMPTION.'),
    ('GAM_W', 'Unit weight of water', 9.81, 'kN/m3', 'Physical constant.'),
    ('FOS_FLOT', 'Required factor of safety against pipe flotation (empty pipe, pipe weight ignored)', 1.1, '-', 'Planning check value.'),
    ('S', 'Pipe-support piles and sheet piles (no pile design available - layout is a planning placeholder)'),
    ('PILE_SP', 'Longitudinal spacing of pile rows under the pipe slab', 3.0, 'm', 'ASSUMPTION - replace with pile design.'),
    ('PILE_THR', 'Outside diameter at or above which two piles per row are used', 1.0, 'm', 'ASSUMPTION.'),
    ('PILE_N_SMALL', 'Piles per row below the threshold', 1, 'no.', 'ASSUMPTION.'),
    ('PILE_N_LARGE', 'Piles per row at or above the threshold', 2, 'no.', 'ASSUMPTION.'),
    ('SHEET_RATIO', 'Sheet pile length / excavation depth from NG (embedment in soft clay)', 2.0, '-', 'ASSUMPTION - temporary works design to confirm.'),
    ('S', 'Programme gates and calendar'),
    ('START', 'Pipe works mobilisation (earliest start of any package)', datetime.date(2026, 12, 1), 'date', 'EXAMPLE - replace with the planned mobilisation date.'),
    ('D_DLU', 'DLU-002 approved / IFC (Rev 03 is Issued for Approval)', datetime.date(2026, 12, 1), 'date', 'Not yet known - currently equal to START (not constraining). Enter the forecast approval date.'),
    ('D_PILEDES', 'Pipe pile and slab design approved', datetime.date(2026, 12, 1), 'date', 'Not yet known - currently equal to START (not constraining). Gates all pipe piling.'),
    ('D_XING', 'Unassigned crossing pipes P0114 / P0117 / P0279 identified', datetime.date(2026, 12, 1), 'date', 'Not yet known - gates excavation of UG-01-S02, UG-02-S04, UG-03-S03.'),
    ('WD', 'Working days per week', 6, 'days', 'Site calendar Sat-Thu (calendar days = working days x 7 / WD).'),
    ('LADDER', 'Overlap within a package: successor starts after this share of its predecessor', 0.3, '-', 'Planning rule (start-to-start ladder). 1.0 = finish-to-start.'),
    ('FFLAG', 'Finish-to-finish lag between consecutive activities in a package', 2, 'cal. days', 'Planning rule.'),
    ('S', 'Productivity per crew / rig (per working day)'),
    ('R_PLAT', 'Set-out, crossing survey and working platform', 80, 'm/day', 'ASSUMPTION.'),
    ('R_PILE', 'Pipe-support piles per rig (600 mm assumed)', 5.0, 'piles/day', "Site actual on aeration 600 mm piling = 5.4 (last 4 full weeks, Aeration sheet). 5.0 used for the longer empty bore from NG."),
    ('R_SHEET', 'Sheet pile installation (vibro / press-in near structures)', 25, 'm wall/day', 'ASSUMPTION.'),
    ('R_DW', 'Dewatering installation along the trench', 40, 'm/day', 'ASSUMPTION.'),
    ('DW_DRAW', 'Drawdown period before excavation', 7, 'days', 'ASSUMPTION - soft silty clay is slow to drain.'),
    ('R_EXS', 'Excavation in a strutted sheet-piled trench', 250, 'm3/day', 'ASSUMPTION.'),
    ('R_EXO', 'Excavation in dewatered open cut', 450, 'm3/day', 'ASSUMPTION.'),
    ('R_SLAB', 'Pile trimming, blinding and piled slab (rebar, form, pour)', 12, 'm/day', 'ASSUMPTION.'),
    ('SLAB_CURE', 'Slab cure before pipe loading', 3, 'days', 'ASSUMPTION.'),
    ('R_PIPE_L', 'Pipe laying, outside diameter >= PIPE_T1', 20, 'm/day', 'ASSUMPTION.'),
    ('R_PIPE_M', 'Pipe laying, PIPE_T2 <= outside diameter < PIPE_T1', 30, 'm/day', 'ASSUMPTION.'),
    ('R_PIPE_S', 'Pipe laying, outside diameter < PIPE_T2', 40, 'm/day', 'ASSUMPTION.'),
    ('PIPE_T1', 'Large-diameter threshold', 2.0, 'm', 'ASSUMPTION.'),
    ('PIPE_T2', 'Medium-diameter threshold', 1.2, 'm', 'ASSUMPTION.'),
    ('FIT_D', 'Extra time per fitting (elbow, tee, reducer)', 0.5, 'days', 'ASSUMPTION.'),
    ('TIE_D', 'Structure connection (puddle flange / grouting) per tie-in package', 3, 'days', 'ASSUMPTION.'),
    ('TEST_D', 'Joint and line test, fixed time per package', 2, 'days', 'ASSUMPTION.'),
    ('R_TEST', 'Surround and testing progress', 150, 'm/day', 'ASSUMPTION.'),
    ('R_BF', 'Backfill and compaction in layers', 350, 'm3/day', 'ASSUMPTION.'),
    ('R_EXT', 'Sheet pile extraction', 50, 'm wall/day', 'ASSUMPTION.'),
]
r = 5
inp_row = {}
for it_ in inputs:
    if it_[0] == 'S':
        wsI.cell(row=r, column=1, value=it_[1]).font = bold
        for c in range(1, 6): wsI.cell(row=r, column=c).fill = sec_fill
        r += 1; continue
    n, d, v, u, s = it_
    put(wsI, r, 1, n, bold); put(wsI, r, 2, d, align=wrap)
    c = put(wsI, r, 3, v, blue, fill=yel)
    if isinstance(v, datetime.date): c.number_format = DATE
    elif isinstance(v, float) and v != int(v): c.number_format = '0.00'
    put(wsI, r, 4, u); put(wsI, r, 5, s, align=wrap)
    name(n, f"Inputs!$C${r}"); inp_row[n] = r
    r += 1
wsI.freeze_panes = 'A5'

# ----------------------------------------------------------------- Aeration (from 3D model)
wsA = wb.create_sheet('Aeration')
wsA['A1'] = 'Aeration piling status and zone forecast (from the 3D model)'; wsA['A1'].font = title
wsA['A2'] = ('Source: Alex-West-Combined-Coordination.blend, scene "AW-AER Piles Actual - 2026-10-18"; tracking data date 17 Oct 2026. '
             'Done = recorded execution, not acceptance. Agreed piling order (18 Oct 2026): Zones 1, 2, 3, then 5 and 4 together, then 6, then 7, then 8.')
wsA['A2'].font = sub
header(wsA, 4, ['Zone', 'Group order', 'Done', 'Pending', 'Defect', 'Unknown', 'Open piles to complete (pending + defect)',
                'Cumulative open piles in agreed order', 'Forecast zone piling complete', 'Pipe packages gated'],
       [8, 10, 9, 9, 9, 9, 16, 16, 16, 14])
cnt = aer.groupby(['zone', 'status']).size().unstack(fill_value=0)
group_of = {1: 1, 2: 2, 3: 3, 4: 4, 5: 4, 6: 5, 7: 6, 8: 7}
zone_row = {}
for i, z in enumerate(range(1, 9)):
    rr = 5 + i; zone_row[z] = rr
    put(wsA, rr, 1, z, bold); put(wsA, rr, 2, group_of[z], blue, fill=yel)
    for j, s in enumerate(['done', 'pending', 'defect', 'unknown']):
        put(wsA, rr, 3 + j, int(cnt.loc[z].get(s, 0)), blue)
    put(wsA, rr, 7, f'=D{rr}+E{rr}')
    put(wsA, rr, 8, f'=SUMIFS($G$5:$G$12,$B$5:$B$12,"<="&B{rr})')
    put(wsA, rr, 9, f'=AER_DATA+ROUNDUP(H{rr}/(AER_RIGS*AER_RATE)*7/WD,0)', fmt=DATE)
    put(wsA, rr, 10, sum(1 for zs in pk.aerz if z in zs))
put(wsA, 13, 1, 'Total', bold)
for j in range(3, 8):
    put(wsA, 13, j, f'=SUM({CL(j)}5:{CL(j)}12)', bold)
wsA['A15'] = 'Production basis'; wsA['A15'].font = bold
rowsA = [('AER_DATA', 'Tracking data date', datetime.date(2026, 10, 17), DATE, 'From model properties (Tracking_Data_Date).'),
         ('AER_RIGS', 'Rigs assumed to stay on aeration piling', 8, '0', 'Model: 8 rigs active in the last two recorded weeks. Edit if rigs move to pipe piling earlier.'),
         ('AER_ACT', 'Actual rate, last 4 full weeks (piles per rig per working day)', None, '0.00', 'Formula from the weekly table below.'),
         ('AER_OVR', 'Override rate (blank = use actual)', None, '0.00', 'Optional.'),
         ('AER_RATE', 'Rate used for the forecast', None, '0.00', '')]
for i, (n, d, v, fmt_, s) in enumerate(rowsA):
    rr = 16 + i
    put(wsA, rr, 1, n, bold); put(wsA, rr, 2, d)
    wsA.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=6)
    if n in ('AER_DATA', 'AER_RIGS'):
        c = put(wsA, rr, 7, v, blue, fmt=fmt_, fill=yel)
    elif n == 'AER_OVR':
        c = put(wsA, rr, 7, None, blue, fmt=fmt_, fill=yel)
    else:
        c = put(wsA, rr, 7, None, fmt=fmt_)
    put(wsA, rr, 8, s); wsA.merge_cells(start_row=rr, start_column=8, end_row=rr, end_column=10)
    name(n, f"Aeration!$G${rr}")
wsA['A23'] = 'Recorded production by week (from execution dates on the pile objects)'; wsA['A23'].font = bold
for j, h in enumerate(['Week from', 'Piles executed', 'Rigs active', 'Piles / rig / working day'], 1):
    c = wsA.cell(row=24, column=j, value=h); c.font = hdr_font; c.fill = hdr_fill; c.alignment = center; c.border = box
for i, w in enumerate(weekly):
    rr = 25 + i
    put(wsA, rr, 1, datetime.date.fromisoformat(w['week']), blue, fmt=DATE)
    put(wsA, rr, 2, w['piles'], blue); put(wsA, rr, 3, w['rigs'], blue)
    put(wsA, rr, 4, f'=B{rr}/C{rr}/WD', fmt='0.00')
last = 25 + len(weekly) - 1
# last 4 full weeks = the four rows before the last (partial) week
wsA['G18'] = f'=SUM(B{last-4}:B{last-1})/SUM(C{last-4}:C{last-1})/WD'
wsA['G20'] = '=IF(AER_OVR="",AER_ACT,AER_OVR)'
wsA[f'E{last}'] = 'Week to data date (17 Oct) - may be partial; excluded from the rate.'; wsA[f'E{last}'].font = sub

# ----------------------------------------------------------------- Structures
wsS = wb.create_sheet('Structures')
wsS['A1'] = 'Structure readiness for pipe tie-ins'; wsS['A1'].font = title
wsS['A2'] = ('Yellow = enter the date each structure is ready for a pipe tie-in (wall and penetration built, its excavation backfilled to pipe zone). '
             'Blank = not constraining. The 3D model holds no construction progress for these structures (only aeration piles), so these dates must come from the master programme.')
wsS['A2'].font = sub
header(wsS, 4, ['Code', 'Structure', 'In 3D model', 'Model stage (manifest 19 Oct 2026)', 'Open gaps (manifest)', 'Ready for tie-in (date)',
                'Tie-in packages', 'Date used'], [10, 40, 11, 34, 60, 15, 10, 12])
srow = {}
codes = sorted(structs, key=lambda c: (c in NOT_MODELLED, c))
for i, c in enumerate(codes):
    rr = 5 + i; srow[c] = rr
    a_ = manifest.get(ASSET_OF.get(c, ''), {})
    put(wsS, rr, 1, c, bold); put(wsS, rr, 2, structs[c])
    put(wsS, rr, 3, 'No' if c in NOT_MODELLED else 'Yes')
    put(wsS, rr, 4, a_.get('stage', 'Not in master - tie-in position not checked' if c in NOT_MODELLED else ''), align=wrap)
    put(wsS, rr, 5, a_.get('gap', ''), align=wrap)
    put(wsS, rr, 6, None, blue, fmt=DATE, fill=yel)
    put(wsS, rr, 7, int(sum(1 for s in pk.struct if c in s)))
    put(wsS, rr, 8, f'=IF(F{rr}="",0,F{rr})', fmt=DATE)
wsS.freeze_panes = 'C5'

# ----------------------------------------------------------------- Fronts
wsF = wb.create_sheet('Fronts')
wsF['A1'] = 'Work fronts and crews'; wsF['A1'].font = title
wsF['A2'] = ('Each front runs its run packages in the order on the Activities sheet (deeper lines first, downstream first). Crew numbers (yellow) divide the durations. '
             'Tie-in packages are outside the chains and assume a separate tie-in team when the structure is ready. Moving a line to another front needs a rebuild.')
wsF['A2'].font = sub
fh = ['Front', 'Description', 'Lines', 'Packages', 'Length (m)', 'Piles (no.)', 'Shored length (m)', 'Piling rigs', 'Sheet pile gangs',
      'Excavation gangs', 'Slab gangs', 'Pipe crews', 'Backfill gangs', 'Start', 'Finish', 'Duration (months)']
header(wsF, 4, fh, [7, 34, 46, 10, 11, 10, 11, 9, 9, 10, 9, 9, 9, 12, 12, 11])
frow = {}
for i, (f, d, ls, _) in enumerate(FRONTS):
    rr = 5 + i; frow[f] = rr
    fl_ = pk[pk.front == f]
    lines_txt = ', '.join(sorted(fl_.line.unique(), key=lambda l: (l[:5], l)))
    put(wsF, rr, 1, f, bold); put(wsF, rr, 2, d); put(wsF, rr, 3, lines_txt, align=wrap)
    ncrew = 2 if sum(plen[p] for p in fl_.id) > 1300 else 1
    put(wsF, rr, 4, f'=COUNTIF(Packages!$C$6:$C$400,A{rr})')
    put(wsF, rr, 5, f'=SUMIF(Packages!$C$6:$C$400,A{rr},Packages!$O$6:$O$400)', fmt='#,##0')
    put(wsF, rr, 6, f'=SUMIF(Packages!$C$6:$C$400,A{rr},Packages!$AC$6:$AC$400)', fmt='#,##0')
    put(wsF, rr, 7, f'=SUMIF(Packages!$C$6:$C$400,A{rr},Packages!$U$6:$U$400)', fmt='#,##0')
    for j in range(8, 14):
        put(wsF, rr, j, ncrew, blue, fill=yel)
    put(wsF, rr, 14, f'=_xlfn.MINIFS(Packages!$AL$6:$AL$400,Packages!$C$6:$C$400,A{rr})', fmt=DATE)
    put(wsF, rr, 15, f'=_xlfn.MAXIFS(Packages!$AM$6:$AM$400,Packages!$C$6:$C$400,A{rr})', fmt=DATE)
    put(wsF, rr, 16, f'=(O{rr}-N{rr})/30.4', fmt='0.0')
tr = 5 + len(FRONTS)
put(wsF, tr, 1, 'All', bold); put(wsF, tr, 2, 'All fronts', bold)
for j in (4, 5, 6, 7):
    put(wsF, tr, j, f'=SUM({CL(j)}5:{CL(j)}{tr-1})', bold, fmt='#,##0')
for j in range(8, 14):
    put(wsF, tr, j, f'=SUM({CL(j)}5:{CL(j)}{tr-1})', bold)
put(wsF, tr, 14, f'=MIN(N5:N{tr-1})', bold, fmt=DATE); put(wsF, tr, 15, f'=MAX(O5:O{tr-1})', bold, fmt=DATE)
put(wsF, tr, 16, f'=(O{tr}-N{tr})/30.4', bold, fmt='0.0')
put(wsF, tr + 1, 2, 'Tie-in packages (separate team)', bold)
put(wsF, tr + 1, 4, '=COUNTIF(Packages!$G$6:$G$400,"Tie-in")+COUNTIF(Packages!$G$6:$G$400,"Short connection")')
FCOL = {'rigs': 'H', 'sheet': 'I', 'exc': 'J', 'slab': 'K', 'pipe': 'L', 'bf': 'M'}

# ----------------------------------------------------------------- Piece check
wsP = wb.create_sheet('Piece check')
wsP['A1'] = 'Piece-level method check with groundwater at GWL'; wsP['A1'].font = title
wsP['A2'] = ('One row per piece from the package workbook (10 m cuts). Open cut only where the dewatered open-cut reach fits short of the nearest other structure or pile '
             'AND the excavation is no deeper than MAXOC below GWL; otherwise sheet piles. Black = formula, blue = source data.')
wsP['A2'].font = sub
ph = ['Line', 'Package', 'Sub-piece', 'Plan length (m)', 'Planning OD (m)', 'Start R (m)', 'End R (m)', 'Start crown (m)', 'End crown (m)',
      'Nearest other structure / pile (m)', 'R deepest (m)', 'R average (m)', 'Excavation depth from NG, avg (m)', 'Depth below GWL, deepest (m)',
      'Depth NG to GWL used, deepest (m)', 'Open-cut reach from axis (m)', 'Method', 'Shored length (m)', 'Depth below GWL, avg (m)',
      'Depth NG to GWL, avg (m)', 'Trench bottom width (m)', 'Excavation shored (m3)', 'Excavation open cut (m3)', 'Piled slab concrete (m3)',
      'Pipe displacement below NG (m3)', 'Flotation FoS start, backfill to NG', 'Flotation FoS end, backfill to NG', 'Flotation FoS start, final to FL',
      'Flotation FoS end, final to FL']
header(wsP, 5, ph, [8, 12, 11, 9, 9, 9, 9, 9, 9, 11, 9, 9, 10, 10, 10, 10, 9, 9, 10, 10, 9, 11, 11, 10, 11, 10, 10, 10, 10])
P0 = 6
def fos(cr, z):
    return (f'=IF({cr}-E{{r}}>=GWL,99,E{{r}}*(MAX(0,GWL-{cr})*(GAM_S-GAM_W)+MAX(0,{z}-MAX({cr},GWL))*GAM_D)/(GAM_W*PI()*E{{r}}^2/4))')
for i, row in pcs.iterrows():
    rr = P0 + i
    vals = [row['Line'], row['Package'], row['Sub-piece ID'], row['Plan length (m)'], row['Planning dia (m)'], row['Start R (m)'],
            row['End R (m)'], row['Start crown (m)'], row['End crown (m)'], row['Nearest other structure or pile from axis (m)']]
    for j, v in enumerate(vals, 1):
        c = wsP.cell(row=rr, column=j, value=v); c.font = blue
    f = {
        11: f'=MIN(F{rr},G{rr})', 12: f'=AVERAGE(F{rr},G{rr})', 13: f'=MAX(0,NG-L{rr})', 14: f'=MAX(0,GWL-K{rr})',
        15: f'=MAX(0,NG-MAX(K{rr},GWL))', 16: f'=E{rr}/2+WS+SLOPE_A*O{rr}+SLOPE_B*N{rr}',
        17: f'=IF(AND(P{rr}<J{rr},N{rr}<=MAXOC),"Open cut","Shoring")', 18: f'=IF(Q{rr}="Shoring",D{rr},0)',
        19: f'=MAX(0,GWL-L{rr})', 20: f'=MAX(0,NG-MAX(L{rr},GWL))', 21: f'=E{rr}+2*WS',
        22: f'=IF(Q{rr}="Shoring",D{rr}*M{rr}*U{rr},0)',
        23: f'=IF(Q{rr}="Open cut",D{rr}*(U{rr}*M{rr}+SLOPE_B*S{rr}^2+2*SLOPE_B*S{rr}*T{rr}+SLOPE_A*T{rr}^2),0)',
        24: f'=IF(M{rr}>0,D{rr}*(E{rr}+2*SLABOVH)*SLAB,0)', 25: f'=IF(M{rr}>0,D{rr}*PI()*E{rr}^2/4,0)',
        26: fos(f'H{rr}', 'NG').format(r=rr), 27: fos(f'I{rr}', 'NG').format(r=rr),
        28: fos(f'H{rr}', 'FL').format(r=rr), 29: fos(f'I{rr}', 'FL').format(r=rr),
    }
    for j, v in f.items():
        c = wsP.cell(row=rr, column=j, value=v); c.font = fnt
for col, fmt_ in zip(range(4, 30), ['0.00'] * 26):
    for rr in range(P0, P0 + len(pcs)):
        wsP.cell(row=rr, column=col).number_format = fmt_ if col not in (17,) else 'General'
PN = P0 + len(pcs) - 1
wsP.freeze_panes = 'D6'
wsP.auto_filter.ref = f'A5:{CL(len(ph))}{PN}'
def PR(col):
    return f"'Piece check'!${col}${P0}:${col}${PN}"

# ----------------------------------------------------------------- Packages
wsK = wb.create_sheet('Packages')
wsK['A1'] = 'Packages - method, quantities, dewatering, flotation and dates'; wsK['A1'].font = title
wsK['A2'] = ('One row per work package from the package workbook (201). Method and quantities recalculate from Inputs and Piece check. '
             'Start / Finish come from the Activities sheet. Pile and sheet quantities are planning placeholders until designs exist.')
wsK['A2'].font = sub
kh = ['Line', 'Package', 'Front', 'Planning name', 'Pipe type', 'Printed size', 'Package type', 'Tie-in structure(s)', 'Structure ready (date)',
      'Aeration zones with open piles within 15 m', 'Aeration gate (date)', 'Crossing ID hold', 'Fittings (no.)', 'Planning OD (m)', 'Plan length (m)',
      'R start (m)', 'R end (m)', 'R deepest (m)', 'Max depth below GWL (m)', 'Max excavation depth from NG (m)', 'Shored length (m)',
      'Open-cut length (m)', 'Excavation method', 'Excavation shored (m3)', 'Excavation open cut (m3)', 'Piled slab concrete (m3)',
      'Pipe displacement (m3)', 'Backfill (m3)', 'Pipe-support piles (no.)', 'Sheet pile wall (m)', 'Sheet pile length (m)', 'Sheet pile area (m2)',
      'Drawdown required below GWL (m)', 'Dewatering method', 'Min flotation FoS, backfill to NG', 'Min flotation FoS, final to FL',
      'Flotation action', 'Start', 'Finish', 'Sequence rank', 'Holds before release (package workbook)', 'Sequence notes (package workbook)',
      'Interactions used in the logic', 'Trench open after laying (days)', 'Dewatering running (days)']
kw = [8, 12, 6, 40, 7, 10, 12, 12, 11, 12, 11, 9, 8, 9, 9, 8, 8, 8, 9, 9, 9, 9, 18, 10, 10, 9, 9, 10, 9, 9, 9, 10, 10, 22, 9, 9, 30, 11, 11, 8, 60, 60, 60, 11, 11]
header(wsK, 5, kh, kw)
K0 = 6
krow = {}
interact_txt = collections.defaultdict(list)
for d_, s_ in deep:
    if d_ in deep_preds[s_]:
        interact_txt[s_].append(f'excavate after {d_} backfilled (deeper)')
        interact_txt[d_].append(f'deeper than {s_}: complete first')
for l_, f_ in shared:
    interact_txt[f_].append(f'shared trench with {l_}: start excavation with it')
    interact_txt[l_].append(f'shared trench with {f_}: backfill after both laid')
for t, nb in closure.items():
    interact_txt[t].append(f'closure: lay after {nb} laid')
for p_, f0_, f1_ in reassigned:
    interact_txt[p_].append(f'built by {f1_} (shared corridor; line system front {f0_})')
for kind_, u_, v_ in dropped:
    if kind_ == 'shared-backfill':
        interact_txt[v_].append(f'COORDINATE ON SITE: do not backfill beside {u_} until it is laid (link left out to keep the logic loop-free)')
order_sorted = sorted(range(len(pk)), key=lambda i: (pk.front[i], pk.grank[i]))
for n_, i in enumerate(order_sorted):
    r_ = pk.iloc[i]; rr = K0 + n_; krow[r_.id] = rr
    put(wsK, rr, 1, r_.line, blue); put(wsK, rr, 2, r_.id, Font(name=F, size=10, bold=True, color='0000FF'))
    put(wsK, rr, 3, r_.front, blue); put(wsK, rr, 4, r_.pname, blue); put(wsK, rr, 5, r_.ptype, blue); put(wsK, rr, 6, r_.size, blue)
    put(wsK, rr, 7, r_.kind, blue); put(wsK, rr, 8, ', '.join(r_.struct), blue)
    if r_.struct:
        put(wsK, rr, 9, '=MAX(' + ','.join(f'Structures!$H${srow[c]}' for c in r_.struct) + ')', green, fmt=DATE)
    else:
        put(wsK, rr, 9, 0, fmt=DATE)
    put(wsK, rr, 10, ', '.join(str(z) for z in r_.aerz) + (f' ({r_.aern} piles)' if r_.aerz else ''), blue)
    if r_.aerz:
        put(wsK, rr, 11, '=MAX(' + ','.join(f'Aeration!$I${zone_row[z]}' for z in r_.aerz) + ')', green, fmt=DATE)
    else:
        put(wsK, rr, 11, 0, fmt=DATE)
    put(wsK, rr, 12, 'Yes' if r_.xing else 'No', blue)
    put(wsK, rr, 13, int(r_.nfit), blue)
    put(wsK, rr, 14, f'=_xlfn.MAXIFS({PR("E")},{PR("B")},B{rr})', green, fmt='0.00')
    put(wsK, rr, 15, f'=SUMIFS({PR("D")},{PR("B")},B{rr})', green, fmt='0.0')
    put(wsK, rr, 16, float(r_.rs), blue, fmt='0.00'); put(wsK, rr, 17, float(r_.re), blue, fmt='0.00')
    put(wsK, rr, 18, f'=_xlfn.MINIFS({PR("K")},{PR("B")},B{rr})', green, fmt='0.00')
    put(wsK, rr, 19, f'=MAX(0,GWL-R{rr})', fmt='0.00')
    put(wsK, rr, 20, f'=MAX(0,NG-R{rr})', fmt='0.00')
    put(wsK, rr, 21, f'=SUMIFS({PR("R")},{PR("B")},B{rr})', green, fmt='0.0')
    put(wsK, rr, 22, f'=O{rr}-U{rr}', fmt='0.0')
    put(wsK, rr, 23, f'=IF(U{rr}<0.05,"Open cut",IF(V{rr}<0.05,"Sheet piles","Mixed: "&TEXT(U{rr},"0")&" m sheet piles"))')
    put(wsK, rr, 24, f'=SUMIFS({PR("V")},{PR("B")},B{rr})', green, fmt='#,##0')
    put(wsK, rr, 25, f'=SUMIFS({PR("W")},{PR("B")},B{rr})', green, fmt='#,##0')
    put(wsK, rr, 26, f'=SUMIFS({PR("X")},{PR("B")},B{rr})', green, fmt='#,##0')
    put(wsK, rr, 27, f'=SUMIFS({PR("Y")},{PR("B")},B{rr})', green, fmt='#,##0')
    put(wsK, rr, 28, f'=MAX(0,X{rr}+Y{rr}-Z{rr}-AA{rr})', fmt='#,##0')
    put(wsK, rr, 29, f'=(ROUNDUP(O{rr}/PILE_SP,0)+1)*IF(N{rr}>=PILE_THR,PILE_N_LARGE,PILE_N_SMALL)', fmt='#,##0')
    put(wsK, rr, 30, f'=2*U{rr}', fmt='#,##0')
    put(wsK, rr, 31, f'=IF(U{rr}>0,ROUNDUP(T{rr}*SHEET_RATIO,0),0)', fmt='0')
    put(wsK, rr, 32, f'=AD{rr}*AE{rr}', fmt='#,##0')
    put(wsK, rr, 33, f'=MAX(0,GWL-(R{rr}-DD_BELOW))', fmt='0.00')
    put(wsK, rr, 34, f'=IF(AG{rr}=0,"Not required",IF(AG{rr}<=WP_LIMIT,"Single-stage wellpoints","Deep wells or two-stage wellpoints"))')
    put(wsK, rr, 35, f'=MIN(_xlfn.MINIFS({PR("Z")},{PR("B")},B{rr}),_xlfn.MINIFS({PR("AA")},{PR("B")},B{rr}))', green, fmt='0.00')
    put(wsK, rr, 36, f'=MIN(_xlfn.MINIFS({PR("AB")},{PR("B")},B{rr}),_xlfn.MINIFS({PR("AC")},{PR("B")},B{rr}))', green, fmt='0.00')
    put(wsK, rr, 37, f'=IF(AJ{rr}<FOS_FLOT,"Anchor pipe to piled slab (permanent straps)",IF(AI{rr}<FOS_FLOT,"Keep dewatering on until fill reaches FL, or strap to slab","No extra measure"))')
    put(wsK, rr, 40, int(r_.grank) + 1, blue)
    put(wsK, rr, 41, r_.holds, blue, align=Alignment(wrap_text=False))
    put(wsK, rr, 42, r_.depends, blue, align=Alignment(wrap_text=False))
    put(wsK, rr, 43, '; '.join(interact_txt[r_.id]) or 'None', blue, align=Alignment(wrap_text=False))
KN = K0 + len(pk) - 1
tot = KN + 1
put(wsK, tot, 2, 'Total', bold)
for col in ['O', 'U', 'V', 'X', 'Y', 'Z', 'AA', 'AB', 'AC', 'AD', 'AF']:
    c = put(wsK, tot, wsK[col + '1'].column, f'=SUM({col}{K0}:{col}{KN})', bold, fmt='#,##0')
put(wsK, tot + 1, 2, 'Check: plan length vs package workbook', sub)
put(wsK, tot + 1, 15, float(seg['Plan length (m)'].sum()), blue, fmt='#,##0.000')
put(wsK, tot + 2, 2, 'Difference (should be 0)', sub)
put(wsK, tot + 2, 15, f'=O{tot}-O{tot+1}', fmt='0.000')
wsK.freeze_panes = 'C6'
wsK.auto_filter.ref = f'A5:{CL(len(kh))}{KN}'

# ----------------------------------------------------------------- Activities
wsT = wb.create_sheet('Activities')
wsT['A1'] = 'Activities and logic (calendar dates; finish = start of next day)'; wsT['A1'].font = title
wsT['A2'] = ('11 activities per package. Start = latest of: predecessor in the package + start-to-start lag (ladder), crew release from the previous package of the same '
             'front (chain), deeper / shared-trench / closure links, and gates. Finish = later of start + duration and predecessor finish + FF lag. Durations recalculate from Inputs and Fronts.')
wsT['A2'].font = sub
th = ['Activity ID', 'WBS', 'Front', 'Line', 'Package', 'Code', 'Activity', 'Quantity', 'Unit', 'Duration (work days)', 'Duration (cal. days)',
      'Ladder predecessor', 'SS lag (cal. days)', 'FF lag (cal. days)', 'Other predecessors (type)', 'Start', 'Finish', 'Crew released',
      'Chain predecessor']
header(wsT, 5, th, [20, 22, 6, 8, 12, 7, 44, 10, 8, 9, 9, 20, 8, 8, 50, 11, 11, 11, 20])
T0 = 6
trow = {}
rowlist = []
for p in [pk.id[i] for i in order_sorted]:
    for c, _ in ACTS:
        rowlist.append((p, c))
for n_, (p, c) in enumerate(rowlist):
    trow[(p, c)] = T0 + n_
QTY = {'A020': ('O', 'm'), 'A030': ('AC', 'piles'), 'A040': ('AD', 'm wall'), 'A050': ('O', 'm'), 'A060': (None, 'm3'),
       'A070': ('O', 'm'), 'A080': ('O', 'm'), 'A090': ('O', 'm'), 'A100': ('AB', 'm3'), 'A110': ('AD', 'm wall')}
for (p, c), rr in trow.items():
    r_ = pk.iloc[idx[p]]; k = krow[p]; fr = frow[r_.front]
    K = lambda col: f'Packages!${col}${k}'
    FR = lambda key_: f'Fronts!${FCOL[key_]}${fr}'
    put(wsT, rr, 1, f'{p}-{c}', blue); put(wsT, rr, 2, f'AW.PIPE.{r_.front}.{r_.line}.{p}', blue)
    put(wsT, rr, 3, r_.front, blue); put(wsT, rr, 4, r_.line, blue); put(wsT, rr, 5, p, blue); put(wsT, rr, 6, c, blue)
    put(wsT, rr, 7, dict(ACTS)[c], blue)
    if c in QTY:
        col, unit = QTY[c]
        put(wsT, rr, 8, f'={K(col)}' if col else f'={K("X")}+{K("Y")}', green, fmt='#,##0')
        put(wsT, rr, 9, unit, blue)
    else:
        put(wsT, rr, 8, None); put(wsT, rr, 9, None)
    dur = {
        'A010': '0',
        'A020': f'MAX(1,ROUNDUP({K("O")}/R_PLAT,0))',
        'A030': f'IF({K("AC")}=0,0,MAX(1,ROUNDUP({K("AC")}/(R_PILE*{FR("rigs")}),0)))',
        'A040': f'IF({K("AD")}=0,0,MAX(1,ROUNDUP({K("AD")}/(R_SHEET*{FR("sheet")}),0)))',
        'A050': f'IF({K("AG")}=0,0,ROUNDUP({K("O")}/R_DW,0)+DW_DRAW)',
        'A060': f'MAX(1,ROUNDUP(({K("X")}/R_EXS+{K("Y")}/R_EXO)/{FR("exc")},0))',
        'A070': f'MAX(1,ROUNDUP({K("O")}/(R_SLAB*{FR("slab")}),0))+SLAB_CURE',
        'A080': (f'MAX(1,ROUNDUP({K("O")}/(IF({K("N")}>=PIPE_T1,R_PIPE_L,IF({K("N")}>=PIPE_T2,R_PIPE_M,R_PIPE_S))*{FR("pipe")})'
                 f'+{K("M")}*FIT_D+IF({K("G")}<>"Run",TIE_D,0),0))'),
        'A090': f'TEST_D+ROUNDUP({K("O")}/R_TEST,0)',
        'A100': f'MAX(1,ROUNDUP({K("AB")}/(R_BF*{FR("bf")}),0))',
        'A110': f'IF({K("AD")}>0,ROUNDUP({K("AD")}/R_EXT,0),0)+IF({K("AG")}>0,1,0)',
    }[c]
    put(wsT, rr, 10, '=' + dur, fmt='0')
    put(wsT, rr, 11, f'=IF(J{rr}=0,0,ROUNDUP(J{rr}*7/WD,0))', fmt='0')
    ci = CODES.index(c)
    others = []
    terms = []
    if c == 'A010':
        put(wsT, rr, 12, None); put(wsT, rr, 13, None); put(wsT, rr, 14, None)
        terms = ['START', 'D_DLU', K('I'), K('K')]
        others.append('gates: START, D_DLU' + (', structure ready' if r_.struct else '') + (', aeration zone piling' if r_.aerz else ''))
    else:
        pr = trow[(p, CODES[ci - 1])]
        put(wsT, rr, 12, f'{p}-{CODES[ci-1]}', blue)
        if c == 'A020':
            put(wsT, rr, 13, None); put(wsT, rr, 14, None)
            terms = [f'Q{pr}']
            others.append('FS from A010')
            if p in chain_prev:
                terms.append(f'P{trow[(chain_prev[p], "A030")]}'); others.append(f'{chain_prev[p]}-A030 (SS, set-out one package ahead of piling)')
            if p in closure:
                terms.append(f'P{trow[(closure[p], "A030")]}'); others.append(f'{closure[p]}-A030 (SS, tie-in team follows the adjacent run)')
        else:
            if c == 'A060':
                put(wsT, rr, 13, f'=IF(K{pr}=0,0,ROUNDUP((K{pr}-ROUNDUP(DW_DRAW*7/WD,0))*LADDER,0)+ROUNDUP(DW_DRAW*7/WD,0))', fmt='0')
            else:
                put(wsT, rr, 13, f'=ROUNDUP(K{pr}*LADDER,0)', fmt='0')
            put(wsT, rr, 14, f'=IF(K{rr}=0,0,FFLAG)', fmt='0')
            terms = [f'P{pr}+M{rr}']
    if c == 'A030':
        terms.append('D_PILEDES'); others.append('gate: D_PILEDES')
    if c == 'A050' and p in chain_prev:
        terms.append(f'Q{trow[(chain_prev[p], "A060")]}-K{rr}'); others.append(f'{chain_prev[p]}-A060 (FS, negative lag = own duration: drawdown complete when the excavation crew is free)')
    if c == 'A060' and r_.xing:
        terms.append('D_XING'); others.append('gate: D_XING')
    chain_ref = None
    if c in CHAINED and p in chain_prev:
        chain_ref = trow[(chain_prev[p], c)]
        terms.append(f'R{chain_ref}')
        put(wsT, rr, 19, f'{chain_prev[p]}-{c}', blue)
    else:
        put(wsT, rr, 19, None)
    if c == 'A060':
        for d_ in deep_preds[p]:
            terms.append(f'Q{trow[(d_, "A100")]}')
        for l_ in ss_preds[p]:
            terms.append(f'P{trow[(l_, "A060")]}')
        others += [f'{d_}-A100 (FS, deeper first)' for d_ in deep_preds[p]] + [f'{l_}-A060 (SS, shared trench)' for l_ in ss_preds[p]]
    if c == 'A080' and p in closure:
        terms.append(f'Q{trow[(closure[p], "A080")]}'); others.append(f'{closure[p]}-A080 (FS, closure)')
    if c == 'A100':
        for o_ in bf_preds[p]:
            terms.append(f'Q{trow[(o_, "A080")]}')
        others += [f'{o_}-A080 (FS, shared trench laid)' for o_ in bf_preds[p]]
    put(wsT, rr, 15, '; '.join(others) or None, blue)
    # columns: P = Start (16), Q = Finish (17), R = Crew released (18)
    put(wsT, rr, 16, '=MAX(' + ','.join(terms) + ')', fmt=DATE)
    if c in ('A010', 'A020'):
        put(wsT, rr, 17, f'=P{rr}+K{rr}', fmt=DATE)
    else:
        put(wsT, rr, 17, f'=MAX(P{rr}+K{rr},Q{pr}+N{rr})', fmt=DATE)
    if c in CHAINED:
        prev_rel = f'R{chain_ref}' if chain_ref else '0'
        put(wsT, rr, 18, f'=IF(K{rr}>0,Q{rr},{prev_rel})', fmt=DATE)
    else:
        put(wsT, rr, 18, None)
TN = T0 + len(rowlist) - 1
wsT.freeze_panes = 'H6'
wsT.auto_filter.ref = f'A5:S{TN}'
# package start/finish on Packages
for p, k in krow.items():
    wsK.cell(row=k, column=38, value=f'=Activities!P{trow[(p, "A020")]}').number_format = DATE
    wsK.cell(row=k, column=39, value=f'=Activities!Q{trow[(p, "A110")]}').number_format = DATE
    wsK.cell(row=k, column=44, value=f'=Activities!P{trow[(p, "A100")]}-Activities!Q{trow[(p, "A080")]}').number_format = '0'
    wsK.cell(row=k, column=45, value=f'=IF(Activities!K{trow[(p, "A050")]}=0,0,Activities!Q{trow[(p, "A110")]}-Activities!P{trow[(p, "A050")]})').number_format = '0'
    for col in (38, 39, 44, 45):
        wsK.cell(row=k, column=col).font = green; wsK.cell(row=k, column=col).border = box

# ----------------------------------------------------------------- Relationships (P6-style)
wsL = wb.create_sheet('Relationships')
wsL['A1'] = 'Relationships (for P6 / MS Project import)'; wsL['A1'].font = title
wsL['A2'] = 'Lags in calendar days, linked to the Activities sheet. Gates (dates) are start constraints on A010 / A030 / A060, listed on the Activities sheet.'; wsL['A2'].font = sub
header(wsL, 4, ['Predecessor', 'Successor', 'Type', 'Lag (cal. days)', 'Reason'], [22, 22, 6, 10, 40])
rl = 5
def rel(pred, succ_, typ, lag, why):
    global rl
    put(wsL, rl, 1, pred, blue); put(wsL, rl, 2, succ_, blue); put(wsL, rl, 3, typ, blue)
    put(wsL, rl, 4, lag, fmt='0'); put(wsL, rl, 5, why, blue); rl += 1
for (p, c), rr in trow.items():
    ci = CODES.index(c)
    if ci == 0:
        continue
    pid = f'{p}-{CODES[ci-1]}'
    if c == 'A020':
        rel(pid, f'{p}-{c}', 'FS', 0, 'Within package')
    else:
        rel(pid, f'{p}-{c}', 'SS', f'=Activities!M{rr}', 'Within package (ladder)')
        rel(pid, f'{p}-{c}', 'FF', f'=Activities!N{rr}', 'Within package (ladder)')
    if c in CHAINED and p in chain_prev:
        rel(f'{chain_prev[p]}-{c}', f'{p}-{c}', 'FS', 0, f'Crew continuity, front {pk.iloc[idx[p]].front}')
    if c == 'A060':
        for d_ in deep_preds[p]: rel(f'{d_}-A100', f'{p}-A060', 'FS', 0, 'Deeper line first')
        for l_ in ss_preds[p]: rel(f'{l_}-A060', f'{p}-A060', 'SS', 0, 'Shared excavation / dewatering front')
    if c == 'A050' and p in chain_prev:
        rel(f'{chain_prev[p]}-A060', f'{p}-A050', 'FS', f'=-Activities!K{rr}', 'Just-in-time dewatering: drawdown complete when excavation crew is free')
    if c == 'A020' and p in chain_prev:
        rel(f'{chain_prev[p]}-A030', f'{p}-A020', 'SS', 0, 'Set-out one package ahead of piling')
    if c == 'A020' and p in closure:
        rel(f'{closure[p]}-A030', f'{p}-A020', 'SS', 0, 'Tie-in team follows the adjacent run')
    if c == 'A080' and p in closure:
        rel(f'{closure[p]}-A080', f'{p}-A080', 'FS', 0, 'Tie-in laid as closure')
    if c == 'A100':
        for o_ in bf_preds[p]: rel(f'{o_}-A080', f'{p}-A100', 'FS', 0, 'Shared trench: all pipes laid before backfill')
wsL.freeze_panes = 'A5'

# ----------------------------------------------------------------- Line summary + Gantt
wsG = wb.create_sheet('Programme')
wsG['A1'] = 'Programme summary and monthly bar chart'; wsG['A1'].font = title
wsG['A2'] = 'Bars recalculate from the Activities sheet. Rows: fronts, then lines. Shaded month = package work active on that line or front.'; wsG['A2'].font = sub
gh = ['Front / line', 'Description', 'Packages', 'Length (m)', 'Shored (m)', 'Piles (no.)', 'Excavation (m3)', 'Start', 'Finish', 'Months']
NM = 24
header(wsG, 5, gh, [12, 34, 9, 10, 10, 9, 11, 11, 11, 8])
for k_ in range(NM):
    c = wsG.cell(row=5, column=11 + k_, value=f'=EDATE(DATE(YEAR(START),MONTH(START),1),{k_})')
    c.number_format = 'mmm-yy'; c.font = hdr_font; c.fill = hdr_fill; c.alignment = Alignment(text_rotation=90, horizontal='center'); c.border = box
    wsG.column_dimensions[CL(11 + k_)].width = 4.2
gr = 6
def grow(label, desc, crit_col, crit, bold_=False):
    global gr
    ft = bold if bold_ else fnt
    put(wsG, gr, 1, label, ft); put(wsG, gr, 2, desc, ft)
    rng = f'Packages!${crit_col}${K0}:${crit_col}${KN}'
    put(wsG, gr, 3, f'=COUNTIF({rng},"{crit}")', ft)
    for j, col in zip(range(4, 8), ['O', 'U', 'AC', None]):
        if col:
            put(wsG, gr, j, f'=SUMIF({rng},"{crit}",Packages!${col}${K0}:${col}${KN})', ft, fmt='#,##0')
        else:
            put(wsG, gr, j, f'=SUMIF({rng},"{crit}",Packages!$X${K0}:$X${KN})+SUMIF({rng},"{crit}",Packages!$Y${K0}:$Y${KN})', ft, fmt='#,##0')
    put(wsG, gr, 8, f'=_xlfn.MINIFS(Packages!$AL${K0}:$AL${KN},{rng},"{crit}")', ft, fmt=DATE)
    put(wsG, gr, 9, f'=_xlfn.MAXIFS(Packages!$AM${K0}:$AM${KN},{rng},"{crit}")', ft, fmt=DATE)
    put(wsG, gr, 10, f'=(I{gr}-H{gr})/30.4', ft, fmt='0.0')
    for k_ in range(NM):
        col = CL(11 + k_)
        c = wsG.cell(row=gr, column=11 + k_, value=f'=IF(AND($H{gr}<=EOMONTH({col}$5,0),$I{gr}>={col}$5),1,"")')
        c.number_format = ';;;'; c.border = box
    gr += 1
put(wsG, gr, 1, 'Fronts', bold); gr += 1
gfront0 = gr
for f, d, ls, _ in FRONTS:
    grow(f, d, 'C', f, True)
gfront1 = gr - 1
put(wsG, gr, 1, 'Lines', bold); gr += 1
for l in lines['Line']:
    lrow = lines[lines.Line == l].iloc[0]
    fts = '/'.join(sorted(pk[pk.line == l].front.unique()))
    grow(l, f"{fts} | {lrow['Printed material']} {lrow['Printed size']} | {lrow['Scheduled from']} to {lrow['Scheduled to']}", 'A', l)
GN = gr - 1
wsG.conditional_formatting.add(f'K{gfront0}:{CL(10+NM)}{gfront1}', CellIsRule(operator='equal', formula=['1'], fill=PatternFill('solid', fgColor='1F3864')))
wsG.conditional_formatting.add(f'K{gfront1+2}:{CL(10+NM)}{GN}', CellIsRule(operator='equal', formula=['1'], fill=PatternFill('solid', fgColor='5B9BD5')))
wsG.freeze_panes = 'C6'
# resource histogram
gr += 1
put(wsG, gr, 1, 'Packages with this activity in the month (sequential packages in a front all count; not concurrent crews)', bold); gr += 1
for code, lab in [('A030', 'Piling'), ('A040', 'Sheet piling'), ('A050', 'Dewatering install/drawdown'), ('A060', 'Excavation'),
                  ('A070', 'Piled slab'), ('A080', 'Pipe laying'), ('A100', 'Backfill')]:
    put(wsG, gr, 1, code, bold); put(wsG, gr, 2, lab)
    for k_ in range(NM):
        col = CL(11 + k_)
        put(wsG, gr, 11 + k_, f'=COUNTIFS(Activities!$F${T0}:$F${TN},"{code}",Activities!$K${T0}:$K${TN},">0",Activities!$P${T0}:$P${TN},"<="&EOMONTH({col}$5,0),Activities!$Q${T0}:$Q${TN},">"&{col}$5)',
            Font(name=F, size=8), fmt='0;;')
    gr += 1
put(wsG, gr, 1, 'Piles', bold); put(wsG, gr, 2, 'Pipe-support piles per month (spread evenly over each piling activity)')
# piles per month = sum over piling activities of qty * overlap days / duration
for k_ in range(NM):
    col = CL(11 + k_)
    put(wsG, gr, 11 + k_, (f'=SUMPRODUCT((Activities!$F${T0}:$F${TN}="A030")*(Activities!$K${T0}:$K${TN}>0),'
                           f'Activities!$H${T0}:$H${TN}*((Activities!$Q${T0}:$Q${TN}<EOMONTH({col}$5,0)+1)*Activities!$Q${T0}:$Q${TN}+(Activities!$Q${T0}:$Q${TN}>=EOMONTH({col}$5,0)+1)*(EOMONTH({col}$5,0)+1)'
                           f'-((Activities!$P${T0}:$P${TN}>{col}$5)*Activities!$P${T0}:$P${TN}+(Activities!$P${T0}:$P${TN}<={col}$5)*{col}$5))'
                           f'*((Activities!$Q${T0}:$Q${TN}>{col}$5)*(Activities!$P${T0}:$P${TN}<EOMONTH({col}$5,0)+1))/(Activities!$K${T0}:$K${TN}+(Activities!$K${T0}:$K${TN}=0)))'),
        Font(name=F, size=8), fmt='0;;')
pile_row = gr
gr += 1
put(wsG, gr, 1, 'Rigs', bold); put(wsG, gr, 2, 'Equivalent piling rigs needed (piles per month / (R_PILE x WD x 52/12))')
for k_ in range(NM):
    put(wsG, gr, 11 + k_, f'={CL(11+k_)}{pile_row}/(R_PILE*WD*52/12)', Font(name=F, size=8, bold=True), fmt='0.0;;')
gr += 1

# ----------------------------------------------------------------- Sensitivity
wsV = wb.create_sheet('Sensitivity')
wsV['A1'] = 'Sheet-pile length under different ground / groundwater rules'; wsV['A1'].font = title
wsV['A2'] = 'Yellow = edit any scenario. Uses the Piece check geometry; the live programme uses the Inputs values (row "Current inputs").'; wsV['A2'].font = sub
header(wsV, 4, ['Scenario', 'Slope above GWL (H:V)', 'Slope below GWL (H:V)', 'Max open-cut depth below GWL (m)', 'Sheet-piled length (m)', 'Share of 8,151 m', 'Note'],
       [34, 12, 12, 14, 14, 11, 70])
scen = [('Package workbook basis (no groundwater)', 1.5, 1.5, 99, 'Reproduces the 1,972 m shoring of the v3 package workbook.'),
        ('Flatter slope below GWL only', 1.5, 2.0, 99, 'Space check only; no depth limit.'),
        ('Depth limit 3.0 m below GWL', 1.5, 2.0, 3.0, ''),
        ('Depth limit 2.5 m below GWL', 1.5, 2.0, 2.5, 'Default in Inputs.'),
        ('Depth limit 2.0 m below GWL', 1.5, 2.0, 2.0, ''),
        ('Everything below GWL sheet-piled', 1.5, 2.0, 0, 'Open cut only where the excavation stays above GWL.')]
for i, (lab, sa, sb, mx, note) in enumerate(scen + [('Current inputs', None, None, None, 'Linked to Inputs.')]):
    rr = 5 + i
    put(wsV, rr, 1, lab, bold if lab == 'Current inputs' else fnt)
    if sa is None:
        put(wsV, rr, 2, '=SLOPE_A', green); put(wsV, rr, 3, '=SLOPE_B', green); put(wsV, rr, 4, '=MAXOC', green)
    else:
        put(wsV, rr, 2, sa, blue, fill=yel); put(wsV, rr, 3, sb, blue, fill=yel); put(wsV, rr, 4, mx, blue, fill=yel)
    put(wsV, rr, 5, (f'=SUMPRODUCT({PR("D")}*(((({PR("E")}/2+WS+B{rr}*{PR("O")}+C{rr}*{PR("N")})>={PR("J")})+({PR("N")}>D{rr}))>0))'), fmt='#,##0')
    put(wsV, rr, 6, f'=E{rr}/SUM({PR("D")})', fmt='0%')
    put(wsV, rr, 7, note)

# ----------------------------------------------------------------- Methodology
wsM = wb.create_sheet('Methodology', 1)
wsM['A1'] = 'Construction methodology - buried GRP / HDPE pipes on piled slabs, groundwater at -1.50'; wsM['A1'].font = title
wsM['A2'] = 'Planning methodology for sequencing and resourcing. Not a temporary-works design or a method statement for approval.'; wsM['A2'].font = sub
header(wsM, 4, ['Step', 'Activity (code)', 'Method', 'Plant / crew', 'Hold points and checks', 'Quantity / duration basis'], [6, 26, 70, 30, 46, 34])
meth = [
    ('1', 'Release (A010)',
     'Package released only when: DLU-002 approved (Rev 03 is Issued for Approval); pile and slab design approved; the connected structure is ready for tie-in packages; '
     'aeration piling in the adjacent zone is complete where open aeration piles are within 15 m; unassigned crossing pipes identified (UG-01-S02, UG-02-S04, UG-03-S03).',
     'Planning / engineering', 'IFC drawing; pile design; permit to dig; structure handover certificate; aeration zone piling complete.',
     'Gates on Inputs, Structures and Aeration sheets.'),
    ('2', 'Set-out and working platform (A020)',
     'Survey pipe axis, slab edges, pile positions and sheet lines. Trial pits / scanning at every crossing listed on Packages. Granular working platform at NG for piling rigs '
     'and cranes on very soft clay; platform designed for rig track pressure.',
     'Survey team, dozer, roller', 'Platform bearing check (plate test) before rig access.', 'Plan length / R_PLAT.'),
    ('3', 'Pipe-support piling (A030)',
     'All pipes are carried on piles and a piled slab. Install piles from the platform at NG BEFORE excavation (CFA or bored, 600 mm assumed), concreting to above cut-off and '
     'leaving the empty-bore length to the slab soffit. Rigs follow the front order; near structures and aeration piles keep the agreed clearances.',
     '1-2 rigs per front (Fronts sheet); site actual 5.4 piles / rig / day, 5.0 used', 'Pile records, integrity tests on every pile, preliminary load test before production.',
     'Piles = (length / PILE_SP + 1) x piles per row - placeholder until pile design.'),
    ('4', 'Sheet piles (A040)',
     'Where the Piece check says sheet piles (open cut does not fit beside a structure / pile, or excavation is deeper than MAXOC below GWL): install interlocking steel sheets '
     'both sides after piling, with a toe in the soft clay for cut-off and base stability. Press-in (silent) method within 10 m of structures and fresh piles to limit vibration.',
     'Sheet pile gang with vibro / press-in rig and crane', 'Temporary works design: embedment, struts, base heave, hydraulic cut-off.',
     'Wall = 2 x shored length; sheet length = depth x SHEET_RATIO.'),
    ('5', 'Dewatering (A050)',
     'Groundwater at -1.50 and excavation bottoms down to -7.6 (199 of 201 packages below GWL, up to 6.1 m) mean almost every package is dewatered. Wellpoints along the trench to 0.5 m below formation; '
     'deep wells or two-stage wellpoints where drawdown exceeds 4.5 m. Draw down for at least DW_DRAW days before excavation. Monitor levels next to existing structures '
     'and fresh piles (settlement / negative skin friction).',
     'Dewatering subcontractor, standby pumps and power', 'Piezometers; settlement markers on structures within 15 m; discharge permit and silt treatment.',
     'Drawdown = GWL - (R - DD_BELOW); method on Packages.'),
    ('6', 'Excavation and strutting (A060)',
     'Excavate in layers, installing walings and struts as each level is reached. Keep the excavator off the edge in soft clay; stockpile away from the trench. Excavate evenly '
     'around the installed piles so they are not pushed sideways. Deeper lines are excavated, laid and backfilled before a shallower line in the same corridor; lines in a '
     'shared trench are opened together.',
     'Long-reach excavator / clamshell, trucks', 'Formation inspection; pile head exposure survey; strut loads.',
     'Shored m3 / R_EXS + open-cut m3 / R_EXO.'),
    ('7', 'Piled slab (A070)',
     'Break down pile heads to cut-off, blinding, reinforcement with pile starter bars, cast the 0.50 m piled slab (pipe cradle or anchor straps cast in where the flotation '
     'check needs them). Cure before pipe load.',
     'Concrete gang, pump', 'Pile head inspection; rebar inspection before pour; cube results.', 'Length / R_SLAB + SLAB_CURE.'),
    ('8', 'Pipe laying and jointing (A080)',
     'Lay from the downstream end upwards (gravity lines). GRP: bell and spigot joints with a crane and lifting beam; HDPE: butt fusion off-line, lowered in strings. Fittings '
     'and thrust restraint as drawn. Tie-in packages are laid last as the closure to the structure, once the structure is ready.',
     'Pipe crew + crane per front; fusion machine for HDPE', 'Joint test per joint (GRP); fusion records (HDPE); alignment and level survey.',
     'Length / pipe rate by OD + fittings x FIT_D + TIE_D at tie-ins.'),
    ('9', 'Surround and testing (A090)',
     'Granular haunch and surround placed and compacted evenly both sides (flexible pipes). Line test by section.', 'Pipe crew', 'Deflection check (GRP); hydrostatic / air test.',
     'TEST_D + length / R_TEST.'),
    ('10', 'Backfill (A100)',
     'Backfill in layers to NG, removing struts progressively as the fill reaches each level. Shared-trench pipes are all laid before backfill starts. Keep dewatering '
     'running until the flotation check is met (Packages column "Flotation action"): empty pipes below -1.50 can float.',
     'Excavator, compactors', 'Compaction tests; flotation check before switching off pumps.', 'Backfill m3 / R_BF.'),
    ('11', 'Dewatering off, extract sheets (A110)',
     'Switch off and remove dewatering only after the fill is heavy enough (or the pipe is strapped to the slab). Extract sheet piles slowly, grouting voids near pipes and '
     'structures. Reinstate the platform / surface.',
     'Crane with vibro extractor', 'Survey pipe level after pumps off; settlement monitoring.', 'Wall m / R_EXT.'),
    ('', 'Sequencing rules (whole network)',
     '1) Eight fronts by system (Fronts sheet). A shared-trench corridor that crosses systems is built by one front (the one holding most of its length; 30 packages moved). '
     '2) Within a front: deeper lines first, downstream first; shared-trench packages back to back along the corridor. 3) Deeper-first links between packages within 8 m with '
     'excavation bottoms more than 1 m apart. 4) Shared trench (bottoms within 1 m): excavate together, backfill only after all its pipes are laid (6 such holds are site '
     'coordination items - see Packages). 5) Set-out runs one package ahead of piling; dewatering is timed so drawdown completes when the excavation crew is free. '
     '6) Tie-ins follow the adjacent run, wait for the structure and close the line. 7) Packages within 15 m of open aeration piles wait for that zone, in the agreed zone order. '
     '8) Pipe piling waits for pile design approval; nothing starts before DLU-002 approval.',
     '', 'Pending decisions: GWL design value confirmation; slope and depth-limit rules; pile layout; structure readiness dates.', ''),
    ('', 'Rig planning note',
     '8 rigs are on aeration piling at about 5.4 piles / rig / working day (measured from the model). 1,587 open aeration piles remain; at that rate they finish around '
     'end November 2026 (Aeration sheet). Pipe-support piling (about 5,800 piles on the placeholder layout) peaks at about 13 equivalent rigs in the first months '
     '(Programme sheet, row Rigs), so releasing the aeration rigs to the pipe fronts is the natural hand-over.',
     '', '', ''),
]
for i, row in enumerate(meth):
    rr = 5 + i
    for j, v in enumerate(row, 1):
        put(wsM, rr, j, v, bold if j <= 2 else fnt, align=wrap)
    wsM.row_dimensions[rr].height = 110 if len(row[2]) > 300 else 80
wsM.freeze_panes = 'A5'

# ----------------------------------------------------------------- Read Me
wsR['A1'] = 'Alex West - underground pipe construction sequence and methodology'; wsR['A1'].font = title
wsR['A2'] = 'Prepared 22 Oct 2026 by Claude for Mohamed. Planning workbook: not a design, temporary-works or construction approval.'; wsR['A2'].font = sub
readme = [
    ('What it is', None),
    (None, 'A sequence, methodology and logic-linked activity programme for the 201 pipe packages (45 lines) of the v3 package workbook, re-checked for groundwater at -1.50 '
           'and with every pipe carried on piles and a piled slab.'),
    ('Inputs used', None),
    (None, '1. Alex-West-Pipe-Execution-Packages-All-Lines.xlsx (v3, 21 Oct 2026): packages, pieces, levels, nearest structures and piles, holds, sequence notes.'),
    (None, '2. Alex-West-Combined-Coordination.blend (19 Oct 2026): aeration pile status for 3,672 piles (data date 17 Oct 2026), agreed aeration piling zone order, '
           'recorded weekly piling production, structure register (stage and open gaps). Check: pile counts near UG-02-S03 reproduce the package workbook (31 within 15 m, 18 pending).'),
    (None, '3. Decisions (Mohamed, 22 Oct 2026): groundwater level -1.50; all pipes need piles and foundations; shoring where needed.'),
    (None, '4. The web viewer link could not be opened from this environment (host blocked); the .blend file was used instead.'),
    ('Sheets', None),
    (None, 'Methodology - step-by-step method, plant, hold points and the sequencing rules.'),
    (None, 'Inputs - levels, groundwater, slopes, pile layout, gates and productivity (yellow = edit).'),
    (None, 'Aeration - pile status by zone from the model, actual production rate, forecast zone completion that gates nearby pipe packages.'),
    (None, 'Structures - enter the date each structure is ready for tie-in (no progress data for these in the model).'),
    (None, 'Fronts - eight work fronts, crew numbers (yellow; 2 crews by default on fronts over 1.3 km) and front start / finish.'),
    (None, 'Packages - method (open cut / sheet piles), excavation, slab, backfill, piles, sheet piles, dewatering, flotation check and dates per package.'),
    (None, 'Activities - 2,211 activities (11 per package) with formula-driven durations and dates; Relationships - the same logic as a P6 / MS Project import list.'),
    (None, 'Programme - front and line summary with a monthly bar chart and monthly counts of activities and piles. Sensitivity - sheet-piled length under other ground rules.'),
    (None, 'Piece check - the piece-level method, volume and flotation calculations behind Packages.'),
    ('How the logic works', None),
    (None, 'Within a package: ladder (start-to-start at LADDER share of the predecessor, finish-to-finish FFLAG). Across packages: crew continuity per front, deeper line first, '
           'shared-trench excavation together, tie-in closure after the adjacent run, gates (DLU approval, pile design, crossing identification, structure readiness, aeration zones).'),
    (None, 'Dates are calendar dates; a finish date is the start of the next free day. Working days convert to calendar days by x 7 / WD.'),
    (None, 'Six shared-trench backfill holds could not be linked without creating a logic loop; they are marked COORDINATE ON SITE in the Packages "Interactions" column.'),
    ('What drives the dates now', None),
    (None, 'START is an example (1 Dec 2026). DLU approval, pile design and crossing-identification dates are set equal to START, so they do not constrain yet. Structure '
           'readiness is blank. Enter these and the programme recalculates. Productivity values are assumptions except pipe piling, which is calibrated to the site aeration piling.'),
    ('Holds that apply to every package', None),
    (None, 'All holds of the package workbook remain (vertical datum not verified, envelope levels, interpolated levels, slab basis, missing pile layouts, structures not in the '
           'master). New: groundwater design value, slopes below GWL and the depth limit for open cut are assumptions for geotechnical confirmation; pipe pile layout is a placeholder.'),
    ('Colour key', None),
    (None, 'Blue text = input or source data; black = formula; green = formula pulling from another sheet; yellow fill = edit.'),
]
rr = 4
for h, t in readme:
    if h:
        wsR.cell(row=rr, column=1, value=h).font = bold
    else:
        c = wsR.cell(row=rr, column=1, value=t); c.font = fnt; c.alignment = Alignment(wrap_text=True, vertical='top')
        wsR.row_dimensions[rr].height = 15 * max(1, len(t) // 150 + 1)
    rr += 1
wsR.column_dimensions['A'].width = 160

for ws in wb.worksheets:
    ws.sheet_view.zoomScale = 90
wb.save(OUT)
print('saved', OUT, 'activities', len(rowlist), 'relationships', rl - 5)
